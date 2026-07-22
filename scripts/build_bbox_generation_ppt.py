from __future__ import annotations

import csv
import json
from collections import defaultdict
from pathlib import Path
from typing import Any

import numpy as np
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

import build_auto_mask_progress_ppt as base


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "reports" / "presentations" / "bbox_generation_progress"
ASSET_DIR = OUT_DIR / "assets"
PPTX_PATH = OUT_DIR / "bbox_generation_bottle_zipper_wood_progress.pptx"
OUTLINE_PATH = OUT_DIR / "bbox_generation_bottle_zipper_wood_progress_outline.md"

BBOX_CSV = ROOT / "reports" / "mvtec_bottle_zipper_qwen_bbox_v2" / "qwen_bbox_eval" / "qwen" / "qwen_bbox_metrics.csv"
BASELINE_CSV = ROOT / "reports" / "mvtec_bottle_zipper_auto_mask" / "official_mask_evaluation" / "auto_mask_metrics.csv"
BBOX_SHEET = ROOT / "reports" / "mvtec_bottle_zipper_qwen_bbox_v2" / "qwen_bbox_eval" / "qwen" / "qwen_bbox_visual_review.png"
BASELINE_SHEET = ROOT / "reports" / "mvtec_bottle_zipper_auto_mask" / "official_mask_evaluation" / "bottle_zipper_auto_mask_visual_review.png"
WOOD_VISUAL = ROOT / "reports" / "auto_mask_phase9_wood" / "auto_masks" / "qwen" / "phase9_wood_visual_result.png"
WOOD_COMPARISON = ROOT / "reports" / "auto_mask_phase9_wood" / "auto_masks" / "qwen" / "visual_result_comparison.png"
WOOD_END_TO_END = ROOT / "reports" / "phase9_reliability_targeted" / "wood_visual_result" / "wood_end_to_end_visual_result.png"
WOOD_CANDIDATES = ROOT / "reports" / "auto_mask_phase9_wood" / "auto_masks" / "qwen" / "contact_sheet_candidate_comparison.png"
ZIPPER_CANDIDATE_MONTAGE = (
    ROOT
    / "outputs"
    / "mvtec_bottle_zipper_qwen_bbox_v2"
    / "qwen_bbox_eval"
    / "qwen"
    / "localization"
    / "zipper"
    / "broken_teeth"
    / "zipper_broken_teeth_002_normal_guided_candidates.png"
)
ZIPPER_HEATMAP = (
    ROOT
    / "outputs"
    / "mvtec_bottle_zipper_qwen_bbox_v2"
    / "qwen_bbox_eval"
    / "qwen"
    / "localization"
    / "zipper"
    / "broken_teeth"
    / "zipper_broken_teeth_002_normal_guided_heatmap.png"
)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    assets = build_assets()
    bbox_rows = load_rows(BBOX_CSV)
    baseline_rows = load_rows(BASELINE_CSV)
    metrics = summarize_bbox_metrics(bbox_rows, baseline_rows)

    prs = Presentation()
    prs.slide_width = base.SLIDE_W
    prs.slide_height = base.SLIDE_H

    add_title_slide(prs, assets)
    add_manual_dependency_slide(prs)
    add_bbox_problem_slide(prs)
    add_architecture_slide(prs)
    add_wood_slide(prs, assets)
    add_bottle_slide(prs, assets, metrics)
    add_zipper_failure_slide(prs, assets, metrics)
    add_normal_guided_slide(prs, assets)
    add_zipper_result_slide(prs, assets, metrics)
    add_design_lessons_slide(prs)
    add_integration_slide(prs)
    add_next_steps_slide(prs)

    set_core_properties(prs)
    prs.save(PPTX_PATH)
    write_outline(metrics)
    validate_presentation(PPTX_PATH, expected_slides=12)
    print(PPTX_PATH)
    print(OUTLINE_PATH)


def set_core_properties(prs: Presentation) -> None:
    props = prs.core_properties
    props.title = "BBox Generation Progress: Wood, Bottle, Zipper"
    props.subject = "ImgGen v2 auto localization and pseudo-mask search-region generation"
    props.author = "ImgGen v2 Research Project"
    props.keywords = "Qwen, bbox, industrial defect, zipper, bottle, wood, PatchCore"
    props.comments = "Generated from current v2 artifacts."


def build_assets() -> dict[str, Path]:
    assets = {
        "wood_visual": copy_or_resize(WOOD_VISUAL, ASSET_DIR / "wood_visual.png", max_width=1600),
        "wood_comparison": copy_or_resize(WOOD_COMPARISON, ASSET_DIR / "wood_comparison.png", max_width=1600),
        "wood_end_to_end_top": crop(WOOD_END_TO_END, ASSET_DIR / "wood_end_to_end_top.png", (0, 0, 1500, 1260)),
        "wood_candidates": copy_or_resize(WOOD_CANDIDATES, ASSET_DIR / "wood_candidates.png", max_width=1800),
        "bbox_bottle": crop_contact_rows(BBOX_SHEET, ASSET_DIR / "bbox_bottle_rows.png", row_start=0, row_count=9),
        "bbox_zipper_teeth": crop_contact_rows(BBOX_SHEET, ASSET_DIR / "bbox_zipper_teeth_rows.png", row_start=9, row_count=3),
        "bbox_zipper_fabric": crop_contact_rows(BBOX_SHEET, ASSET_DIR / "bbox_zipper_fabric_rows.png", row_start=12, row_count=3),
        "bbox_zipper_split": crop_contact_rows(BBOX_SHEET, ASSET_DIR / "bbox_zipper_split_rows.png", row_start=15, row_count=3),
        "baseline_zipper": crop_contact_rows(BASELINE_SHEET, ASSET_DIR / "baseline_zipper_rows.png", row_start=9, row_count=9, thumb=180, label=42, header=40),
        "zipper_candidates": copy_or_resize(ZIPPER_CANDIDATE_MONTAGE, ASSET_DIR / "zipper_candidate_montage.png", max_width=1200),
        "zipper_heatmap": colorize_heatmap(ZIPPER_HEATMAP, ASSET_DIR / "zipper_heatmap_color.png"),
    }
    return assets


def crop_contact_rows(
    src: Path,
    dst: Path,
    *,
    row_start: int,
    row_count: int,
    thumb: int = 190,
    label: int = 50,
    header: int = 38,
) -> Path:
    image = Image.open(src).convert("RGB")
    row_h = thumb + label
    top = max(0, header + row_start * row_h)
    bottom = min(image.height, header + (row_start + row_count) * row_h)
    section = image.crop((0, top, image.width, bottom))
    header_im = image.crop((0, 0, image.width, header))
    result = Image.new("RGB", (image.width, header + section.height), "white")
    result.paste(header_im, (0, 0))
    result.paste(section, (0, header))
    dst.parent.mkdir(parents=True, exist_ok=True)
    result.save(dst)
    return dst


def crop(src: Path, dst: Path, box: tuple[int, int, int, int]) -> Path:
    image = Image.open(src).convert("RGB")
    box = (max(0, box[0]), max(0, box[1]), min(image.width, box[2]), min(image.height, box[3]))
    result = image.crop(box)
    dst.parent.mkdir(parents=True, exist_ok=True)
    result.save(dst)
    return dst


def copy_or_resize(src: Path, dst: Path, *, max_width: int) -> Path:
    image = Image.open(src).convert("RGB")
    if image.width > max_width:
        scale = max_width / image.width
        image = image.resize((max_width, int(round(image.height * scale))), Image.Resampling.LANCZOS)
    dst.parent.mkdir(parents=True, exist_ok=True)
    image.save(dst)
    return dst


def colorize_heatmap(src: Path, dst: Path) -> Path:
    gray = Image.open(src).convert("L")
    arr = np.asarray(gray, dtype=np.float32) / 255.0
    rgb = np.zeros((arr.shape[0], arr.shape[1], 3), dtype=np.uint8)
    rgb[..., 0] = np.uint8(np.clip((arr - 0.25) / 0.75, 0, 1) * 255)
    rgb[..., 1] = np.uint8(np.clip(1.0 - np.abs(arr - 0.5) * 2.0, 0, 1) * 210)
    rgb[..., 2] = np.uint8(np.clip((0.75 - arr) / 0.75, 0, 1) * 255)
    image = Image.fromarray(rgb, mode="RGB")
    dst.parent.mkdir(parents=True, exist_ok=True)
    image.save(dst)
    return dst


def load_rows(path: Path) -> list[dict[str, Any]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def summarize_bbox_metrics(new_rows: list[dict[str, Any]], old_rows: list[dict[str, Any]]) -> dict[str, dict[str, float]]:
    old_by_key = {(row["category"], row["defect_type"]): [] for row in old_rows}
    for row in old_rows:
        old_by_key.setdefault((row["category"], row["defect_type"]), []).append(row)
    new_by_key = {(row["category"], row["defect_type"]): [] for row in new_rows}
    for row in new_rows:
        new_by_key.setdefault((row["category"], row["defect_type"]), []).append(row)
    keys = sorted(set(old_by_key) | set(new_by_key))
    result: dict[str, dict[str, float]] = {}
    for key in keys:
        old = old_by_key.get(key, [])
        new = new_by_key.get(key, [])
        result[f"{key[0]}/{key[1]}"] = {
            "old_recall": mean(old, "qwen_region_recall"),
            "new_recall": mean(new, "qwen_region_recall"),
            "old_iou": mean(old, "qwen_region_iou"),
            "new_iou": mean(new, "qwen_region_iou"),
            "new_area": mean(new, "qwen_region_area_fraction"),
        }
    for category in sorted({row["category"] for row in new_rows}):
        old = [row for row in old_rows if row["category"] == category]
        new = [row for row in new_rows if row["category"] == category]
        result[f"{category}_overall"] = {
            "old_recall": mean(old, "qwen_region_recall"),
            "new_recall": mean(new, "qwen_region_recall"),
            "old_iou": mean(old, "qwen_region_iou"),
            "new_iou": mean(new, "qwen_region_iou"),
            "new_area": mean(new, "qwen_region_area_fraction"),
        }
    return result


def mean(rows: list[dict[str, Any]], key: str) -> float:
    values = []
    for row in rows:
        try:
            values.append(float(row[key]))
        except Exception:
            pass
    return float(np.mean(values)) if values else 0.0


def add_title_slide(prs: Presentation, assets: dict[str, Path]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = base.NAVY
    base.add_rect(slide, 0, 0, 0.18, 7.5, base.TEAL, line=None)
    base.add_text(slide, 0.72, 0.72, 11.5, 0.5, "BBOX GENERATION PROGRESS", 13, base.CYAN, bold=True)
    base.add_text(slide, 0.72, 1.28, 7.5, 1.22, "Automatic localization for\nwood, bottle, and zipper defects", 31, base.WHITE, bold=True)
    base.add_text(slide, 0.75, 2.92, 7.05, 0.72, "Current upgrade: Qwen semantic localization plus normal-guided candidate verification for dense zipper texture.", 16, base.PALE)
    base.add_pill(slide, 0.75, 4.06, 1.42, 0.38, "Wood", base.ORANGE)
    base.add_pill(slide, 2.32, 4.06, 1.42, 0.38, "Bottle", base.TEAL)
    base.add_pill(slide, 3.89, 4.06, 1.42, 0.38, "Zipper", base.PURPLE)
    base.add_pill(slide, 5.46, 4.06, 2.1, 0.38, "High-recall bbox", base.GREEN)
    base.add_text(slide, 0.75, 6.54, 6.8, 0.3, "ImgGen v2 current progress | June 2026", 11, base.MUTED)
    base.add_picture_cover(slide, assets["wood_comparison"], 8.05, 0.82, 4.75, 5.55, radius=True)


def add_manual_dependency_slide(prs: Presentation) -> None:
    slide = base.blank_slide(prs, "Why bbox generation matters", "Motivation")
    base.add_text(slide, 0.68, 0.92, 12.0, 0.55, "Industrial defect generation usually needs spatial control.", 22, base.NAVY, bold=True, align=PP_ALIGN.CENTER)
    base.add_text(slide, 1.0, 1.54, 11.3, 0.32, "In practice, many systems start from a manually supplied bounding box or mask.", 13, base.SLATE, align=PP_ALIGN.CENTER)
    stages = [
        ("Human", "inspect image", base.ORANGE),
        ("Draw", "bbox / mask", base.ORANGE),
        ("Generate", "defect image", base.PURPLE),
        ("Train", "detector", base.TEAL),
    ]
    x = 1.05
    for i, (title, body, color) in enumerate(stages):
        base.draw_pipeline_box(slide, x + i * 3.0, 2.28, 2.35, 1.18, title, body, color, str(i + 1))
        if i < len(stages) - 1:
            base.add_arrow(slide, x + i * 3.0 + 2.38, 2.88, x + i * 3.0 + 2.74, 2.88, base.SLATE)
    base.add_callout(slide, 0.85, 4.15, 3.55, 1.2, "Cost", "Manual masks are slow; boxes are cheaper but still human-dependent.", base.ORANGE)
    base.add_callout(slide, 4.88, 4.15, 3.55, 1.2, "Noise", "Rough boxes can include normal texture or miss weak defect regions.", base.RED)
    base.add_callout(slide, 8.91, 4.15, 3.55, 1.2, "Scale", "Every new part and defect type repeats the annotation loop.", base.PURPLE)
    base.add_text(slide, 0.85, 6.28, 11.7, 0.42, "Goal: generate high-recall search regions automatically from normal images, defect examples, and optional text.", 15, base.TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)


def add_bbox_problem_slide(prs: Presentation) -> None:
    slide = base.blank_slide(prs, "BBox objective: high recall first", "Principle")
    base.add_text(slide, 0.72, 0.95, 12.0, 0.38, "A bbox is not the final mask. It is the search region for downstream pseudo-mask refinement.", 17, base.NAVY, bold=True, align=PP_ALIGN.CENTER)
    base.add_rect(slide, 0.9, 1.65, 5.65, 3.65, base.WHITE, line=base.PALE, radius=True)
    base.add_text(slide, 1.2, 1.95, 5.05, 0.38, "Bad bbox failure", 18, base.RED, bold=True, align=PP_ALIGN.CENTER)
    base.add_bullet_list(slide, 1.25, 2.55, 4.9, 1.55, ["Misses the true defect", "Downstream masks cannot recover excluded pixels", "Looks precise but fails recall"], 13, base.NAVY, bullet_color=base.RED)
    base.add_text(slide, 1.2, 4.58, 5.05, 0.4, "Fatal for auto-mask generation", 15, base.RED, bold=True, align=PP_ALIGN.CENTER)
    base.add_rect(slide, 6.85, 1.65, 5.65, 3.65, base.WHITE, line=base.PALE, radius=True)
    base.add_text(slide, 7.15, 1.95, 5.05, 0.38, "Acceptable bbox tradeoff", 18, base.GREEN, bold=True, align=PP_ALIGN.CENTER)
    base.add_bullet_list(slide, 7.2, 2.55, 4.9, 1.55, ["Contains the full defect", "May include some normal material", "Refinement can remove background"], 13, base.NAVY, bullet_color=base.GREEN)
    base.add_text(slide, 7.15, 4.58, 5.05, 0.4, "Correct for search-region generation", 15, base.GREEN, bold=True, align=PP_ALIGN.CENTER)
    base.add_text(slide, 0.85, 6.15, 11.7, 0.48, "Professional rule: optimize bbox recall before bbox tightness. Tight boxes are useful only when they still contain the defect.", 15, base.TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)


def add_architecture_slide(prs: Presentation) -> None:
    slide = base.blank_slide(prs, "BBox generation architecture", "System")
    stages = [
        ("Inputs", "normal images\nfew defect images\noptional text", base.GREEN),
        ("Qwen", "semantic defect cue\nrough region", base.CYAN),
        ("Normal memory", "PatchCore-style\nanomaly heatmap", base.PURPLE),
        ("Verifier", "Qwen chooses\ncandidate windows", base.ORANGE),
        ("BBox output", "high-recall\nsearch region", base.TEAL),
    ]
    for i, (title, body, color) in enumerate(stages):
        x = 0.55 + i * 2.53
        base.draw_stage_card(slide, x, 1.25, 2.18, 1.55, title, body, color, i + 1)
        if i < len(stages) - 1:
            base.add_arrow(slide, x + 2.2, 2.03, x + 2.45, 2.03, base.SLATE)
    base.add_text(slide, 0.7, 3.35, 5.65, 0.35, "Original mode: Qwen bbox", 17, base.NAVY, bold=True)
    base.add_bullet_list(slide, 0.82, 3.85, 5.65, 1.55, ["Works reasonably for visible bottle regions", "Works for wood when used as broad search fence", "Fails on zipper teeth because repeated texture confuses spatial location"], 12, base.NAVY, bullet_color=base.ORANGE)
    base.add_text(slide, 7.05, 3.35, 5.65, 0.35, "New mode: normal_guided_qwen_verify", 17, base.NAVY, bold=True)
    base.add_bullet_list(slide, 7.17, 3.85, 5.65, 1.55, ["Normal memory proposes candidate y-windows", "Qwen verifies semantic match among candidates", "Search region becomes broad but high-recall"], 12, base.NAVY, bullet_color=base.TEAL)
    base.add_text(slide, 0.85, 6.35, 11.7, 0.3, "This is still bbox generation. Pixel masks are generated later from candidate/refinement evidence.", 12, base.MUTED, align=PP_ALIGN.CENTER)


def add_wood_slide(prs: Presentation, assets: dict[str, Path]) -> None:
    slide = base.blank_slide(prs, "Wood: Qwen as a search fence plus morphology policy", "Wood")
    base.add_picture_contain(slide, assets["wood_visual"], 0.45, 1.05, 7.65, 5.7)
    base.add_rect(slide, 8.35, 1.05, 4.4, 5.7, base.WHITE, line=base.PALE, radius=True)
    base.add_text(slide, 8.65, 1.35, 3.8, 0.35, "Current wood routing", 18, base.NAVY, bold=True, align=PP_ALIGN.CENTER)
    rows = [
        ("000", "multi_scuff", "patchcore_guided", "training_soft", base.ORANGE),
        ("001", "scratch_band", "ensemble_consensus", "training_medium", base.GREEN),
        ("002", "scratch_band", "ensemble_consensus", "training_medium", base.GREEN),
    ]
    y = 1.95
    for sample, morph, selected, target, color in rows:
        base.add_result_row(slide, 8.65, y, 3.75, 0.9, sample, selected, morph, target, color, compact=True)
        y += 1.08
    base.add_callout(slide, 8.65, 5.48, 3.75, 0.78, "Takeaway", "Wood bbox can be broad because the candidate bank and morphology policy determine the pseudo-mask.", base.TEAL)


def add_bottle_slide(prs: Presentation, assets: dict[str, Path], metrics: dict[str, dict[str, float]]) -> None:
    slide = base.blank_slide(prs, "Bottle: visible defects are mostly localizable by Qwen", "Bottle")
    base.add_picture_contain(slide, assets["bbox_bottle"], 0.42, 1.02, 7.8, 5.95)
    base.add_text(slide, 8.55, 1.05, 3.95, 0.35, "Old Qwen recall", 18, base.NAVY, bold=True, align=PP_ALIGN.CENTER)
    bottle = metrics["bottle_overall"]
    base.add_metric_card(slide, 8.75, 1.65, 3.45, 1.35, f"{bottle['old_recall']:.3f}", "overall recall", "baseline prompt", base.TEAL)
    base.add_metric_card(slide, 8.75, 3.25, 3.45, 1.35, f"{bottle['new_recall']:.3f}", "v2 tight recall", "stricter bbox prompt", base.ORANGE)
    base.add_text(slide, 8.52, 4.98, 4.05, 0.88, "Stricter prompt made some boxes tighter but did not improve bottle aggregate recall. For bottle, keep Qwen as a high-recall coarse search fence.", 12, base.NAVY, align=PP_ALIGN.CENTER)
    add_metric_table(slide, 8.45, 6.0, [
        ("broken_large", metrics["bottle/broken_large"]),
        ("broken_small", metrics["bottle/broken_small"]),
        ("contamination", metrics["bottle/contamination"]),
    ])


def add_zipper_failure_slide(prs: Presentation, assets: dict[str, Path], metrics: dict[str, dict[str, float]]) -> None:
    slide = base.blank_slide(prs, "Zipper: direct Qwen bbox failed on dense repeated teeth", "Zipper")
    base.add_picture_contain(slide, assets["baseline_zipper"], 0.45, 1.02, 7.0, 5.95)
    base.add_rect(slide, 7.95, 1.05, 4.75, 5.8, base.WHITE, line=base.PALE, radius=True)
    base.add_text(slide, 8.25, 1.35, 4.15, 0.35, "Baseline failure mode", 18, base.RED, bold=True, align=PP_ALIGN.CENTER)
    base.add_bullet_list(slide, 8.35, 1.98, 4.05, 1.28, ["Qwen selected full vertical chain or wrong tooth", "Broken teeth recall was nearly zero", "Split teeth was inconsistent"], 12, base.NAVY, bullet_color=base.RED)
    base.add_metric_card(slide, 8.35, 3.65, 1.9, 1.28, f"{metrics['zipper/broken_teeth']['old_recall']:.3f}", "broken_teeth", "old recall", base.RED)
    base.add_metric_card(slide, 10.55, 3.65, 1.9, 1.28, f"{metrics['zipper/split_teeth']['old_recall']:.3f}", "split_teeth", "old recall", base.RED)
    base.add_text(slide, 8.35, 5.42, 4.05, 0.78, "Root cause: zipper defects compete with dense normal repeated teeth, so visual-language coordinate drawing alone is unstable.", 12, base.NAVY, bold=True, align=PP_ALIGN.CENTER)


def add_normal_guided_slide(prs: Presentation, assets: dict[str, Path]) -> None:
    slide = base.blank_slide(prs, "Upgrade: normal-guided Qwen verification", "Zipper")
    base.add_picture_cover(slide, assets["zipper_heatmap"], 0.7, 1.18, 2.55, 4.45, radius=True)
    base.add_arrow(slide, 3.35, 3.4, 3.75, 3.4, base.SLATE)
    base.add_picture_contain(slide, assets["zipper_candidates"], 3.95, 1.18, 4.9, 4.45)
    base.add_arrow(slide, 8.95, 3.4, 9.35, 3.4, base.SLATE)
    base.add_rect(slide, 9.6, 1.18, 2.95, 4.45, base.WHITE, line=base.PALE, radius=True)
    base.add_text(slide, 9.85, 1.52, 2.45, 0.3, "Qwen verifier", 17, base.PURPLE, bold=True, align=PP_ALIGN.CENTER)
    base.add_bullet_list(slide, 9.92, 2.18, 2.25, 1.7, ["Receives numbered candidate crops", "Chooses up to 3 likely defect windows", "Output bbox is union of selected windows"], 11, base.NAVY, bullet_color=base.PURPLE)
    base.add_pill(slide, 10.1, 4.45, 2.05, 0.4, "high-recall bbox", base.GREEN, font_size=10)
    base.add_text(slide, 0.85, 5.95, 11.7, 0.55, "Key shift: Qwen stops drawing exact coordinates and becomes a semantic judge over normal-memory candidates.", 15, base.TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)


def add_zipper_result_slide(prs: Presentation, assets: dict[str, Path], metrics: dict[str, dict[str, float]]) -> None:
    slide = base.blank_slide(prs, "Zipper result: recall recovered", "Results")
    base.add_picture_contain(slide, assets["bbox_zipper_teeth"], 0.45, 1.0, 3.8, 5.75)
    base.add_picture_contain(slide, assets["bbox_zipper_split"], 4.45, 1.0, 3.8, 5.75)
    zipper = metrics["zipper_overall"]
    bt = metrics["zipper/broken_teeth"]
    st = metrics["zipper/split_teeth"]
    base.add_metric_card(slide, 8.72, 1.15, 3.35, 1.25, f"{bt['old_recall']:.3f} → {bt['new_recall']:.3f}", "broken_teeth recall", "normal-guided verifier", base.GREEN)
    base.add_metric_card(slide, 8.72, 2.75, 3.35, 1.25, f"{st['old_recall']:.3f} → {st['new_recall']:.3f}", "split_teeth recall", "normal-guided verifier", base.GREEN)
    base.add_metric_card(slide, 8.72, 4.35, 3.35, 1.25, f"{zipper['old_recall']:.3f} → {zipper['new_recall']:.3f}", "zipper overall recall", "all selected defects", base.TEAL)
    base.add_text(slide, 8.45, 6.0, 3.9, 0.58, "IoU also improved, but boxes remain broad by design. This is correct for search-region generation.", 12, base.NAVY, bold=True, align=PP_ALIGN.CENTER)


def add_design_lessons_slide(prs: Presentation) -> None:
    slide = base.blank_slide(prs, "Professional lessons from wood, bottle, and zipper", "Lessons")
    lessons = [
        ("Wood", "Qwen bbox is enough as a search fence; morphology policy decides hard vs soft labels.", base.ORANGE),
        ("Bottle", "Visible object-level defects are mostly Qwen-localizable, but over-tight prompting can reduce recall.", base.TEAL),
        ("Zipper", "Dense repeated texture needs normal-guided proposals before Qwen verification.", base.PURPLE),
    ]
    for i, (title, body, color) in enumerate(lessons):
        base.add_callout(slide, 0.85, 1.15 + i * 1.55, 11.6, 1.05, title, body, color)
    base.add_rect(slide, 1.1, 6.05, 11.15, 0.75, base.NAVY, line=None, radius=True)
    base.add_text(slide, 1.38, 6.27, 10.55, 0.32, "One bbox policy is not enough. Use category/defect-specific localization policies.", 15, base.WHITE, bold=True, align=PP_ALIGN.CENTER)


def add_integration_slide(prs: Presentation) -> None:
    slide = base.blank_slide(prs, "Pipeline integration", "Implementation")
    base.add_text(slide, 0.75, 1.0, 12, 0.35, "The new policy is opt-in per target, so bottle and wood behavior stays stable.", 16, base.NAVY, bold=True, align=PP_ALIGN.CENTER)
    code = (
        "auto_masks:\n"
        "  localization_policy_by_target:\n"
        "    zipper/broken_teeth: normal_guided_qwen_verify\n"
        "    zipper/split_teeth: normal_guided_qwen_verify\n"
        "  normal_guided_window_height_px: 176\n"
        "  normal_guided_candidate_count: 6\n"
        "  normal_guided_max_selected_candidates: 3"
    )
    base.add_rect(slide, 0.95, 1.72, 5.8, 3.0, base.NAVY, line=None, radius=True)
    base.add_text(slide, 1.25, 2.0, 5.25, 2.4, code, 15, base.WHITE, font=base.FONT_MONO)
    base.add_text(slide, 7.25, 1.78, 4.9, 0.35, "Artifacts produced", 18, base.NAVY, bold=True)
    base.add_bullet_list(slide, 7.25, 2.35, 4.8, 1.65, ["bbox metrics CSV", "raw Qwen JSONL", "candidate montages", "normal-guided heatmaps", "visual before/after sheet"], 13, base.NAVY, bullet_color=base.TEAL)
    base.add_text(slide, 7.25, 4.6, 4.8, 0.55, "Safety: if normal-guided verification fails, auto-mask falls back to the original Qwen bbox path.", 12, base.RED, bold=True, align=PP_ALIGN.CENTER)
    base.add_text(slide, 1.2, 5.65, 11.0, 0.44, "Implemented in iadgen_v2/auto_masks.py and evaluated by scripts/evaluate_qwen_bbox_prompts.py", 12, base.MUTED, align=PP_ALIGN.CENTER)


def add_next_steps_slide(prs: Presentation) -> None:
    slide = base.blank_slide(prs, "Next steps", "Roadmap")
    steps = [
        ("1", "Rerun full auto-mask on bottle/zipper", "Measure whether high-recall zipper bboxes improve selected masks, not only search-region recall.", base.TEAL),
        ("2", "Tighten normal-guided windows", "Tune x-range and candidate union so recall stays high while area fraction drops.", base.ORANGE),
        ("3", "Extend policy routing", "Use normal_guided_qwen_verify for other dense repeated textures after validation.", base.PURPLE),
        ("4", "Cache normal-memory features", "Reduce runtime for repeated bbox experiments and full Phase 9 runs.", base.GREEN),
    ]
    for i, (num, title, body, color) in enumerate(steps):
        x = 0.75 + (i % 2) * 6.15
        y = 1.15 + (i // 2) * 2.4
        base.draw_pipeline_box(slide, x, y, 5.25, 1.6, title, body, color, num)
    base.add_text(slide, 0.85, 6.35, 11.7, 0.35, "Immediate validation target: rerun the official bottle/zipper auto-mask evaluation with the new localization policy enabled.", 13, base.NAVY, bold=True, align=PP_ALIGN.CENTER)


def add_metric_table(slide: object, x: float, y: float, rows: list[tuple[str, dict[str, float]]]) -> None:
    columns = ["Defect", "old R", "new R"]
    table_rows = [[name, f"{vals['old_recall']:.3f}", f"{vals['new_recall']:.3f}"] for name, vals in rows]
    base.add_table(slide, x, y, 4.05, 0.75, columns, table_rows, [1.55, 0.82, 0.82])


def write_outline(metrics: dict[str, dict[str, float]]) -> None:
    lines = [
        "# BBox Generation Progress PPT Outline",
        "",
        f"PPTX: `{PPTX_PATH}`",
        "",
        "## Main Claim",
        "",
        "BBox generation should prioritize high-recall search regions. Wood and bottle can use Qwen as a coarse semantic locator, while zipper teeth require normal-memory candidate proposals plus Qwen verification.",
        "",
        "## Key Result",
        "",
        f"- Zipper overall recall improved from `{metrics['zipper_overall']['old_recall']:.4f}` to `{metrics['zipper_overall']['new_recall']:.4f}`.",
        f"- Broken teeth recall improved from `{metrics['zipper/broken_teeth']['old_recall']:.4f}` to `{metrics['zipper/broken_teeth']['new_recall']:.4f}`.",
        f"- Split teeth recall improved from `{metrics['zipper/split_teeth']['old_recall']:.4f}` to `{metrics['zipper/split_teeth']['new_recall']:.4f}`.",
        "",
        "## Slides",
        "",
        "1. Title",
        "2. Why bbox generation matters",
        "3. BBox objective: high recall first",
        "4. Architecture",
        "5. Wood case",
        "6. Bottle case",
        "7. Zipper failure",
        "8. Normal-guided Qwen verification",
        "9. Zipper result",
        "10. Lessons",
        "11. Pipeline integration",
        "12. Next steps",
    ]
    OUTLINE_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")


def validate_presentation(path: Path, *, expected_slides: int) -> None:
    prs = Presentation(path)
    if len(prs.slides) != expected_slides:
        raise RuntimeError(f"Expected {expected_slides} slides, found {len(prs.slides)}")
    if path.stat().st_size < 200_000:
        raise RuntimeError(f"Generated presentation looks too small: {path.stat().st_size} bytes")


if __name__ == "__main__":
    main()
