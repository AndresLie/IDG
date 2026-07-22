from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Iterable

import numpy as np
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "reports" / "presentations" / "auto_mask_generation_progress"
ASSET_DIR = OUT_DIR / "assets"
PPTX_PATH = OUT_DIR / "auto_mask_generation_current_progress.pptx"
OUTLINE_PATH = OUT_DIR / "auto_mask_generation_current_progress_outline.md"

DATA_ROOT = ROOT / "data" / "auto_mask_smoke" / "custom_part"
MASK_ROOT = ROOT / "outputs" / "auto_mask_phase9_wood" / "auto_masks" / "qwen"
REPORT_ROOT = ROOT / "reports" / "auto_mask_phase9_wood"
TARGETED_REPORT = ROOT / "reports" / "phase9_reliability_targeted"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

NAVY = RGBColor(15, 23, 42)
NAVY_2 = RGBColor(30, 41, 59)
SLATE = RGBColor(51, 65, 85)
MUTED = RGBColor(100, 116, 139)
LIGHT = RGBColor(248, 250, 252)
PALE = RGBColor(241, 245, 249)
WHITE = RGBColor(255, 255, 255)
TEAL = RGBColor(13, 148, 136)
TEAL_DARK = RGBColor(15, 118, 110)
CYAN = RGBColor(14, 165, 233)
ORANGE = RGBColor(245, 158, 11)
RED = RGBColor(239, 68, 68)
GREEN = RGBColor(34, 197, 94)
PURPLE = RGBColor(139, 92, 246)

FONT = "Aptos"
FONT_MONO = "Aptos Mono"


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    assets = build_assets()
    records = load_auto_mask_records()
    ablation = load_ablation_rows()

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    add_title_slide(prs, assets)
    add_manual_dependency_slide(prs)
    add_research_gap_slide(prs)
    add_progress_snapshot_slide(prs)
    add_architecture_slide(prs)
    add_input_contract_slide(prs, assets)
    add_qwen_localization_slide(prs, assets, records)
    add_candidate_bank_slide(prs)
    add_candidate_scoring_slide(prs)
    add_scuff_case_slide(prs, assets, records)
    add_scratch_cases_slide(prs, assets, records)
    add_policy_routing_slide(prs)
    add_mask_variants_slide(prs, assets)
    add_uncertainty_slide(prs, assets)
    add_wood_results_slide(prs, assets, records)
    add_candidate_comparison_slide(prs, assets)
    add_ablation_slide(prs, assets, ablation)
    add_manifest_validation_slide(prs)
    add_downstream_integration_slide(prs, assets)
    add_strengths_limitations_slide(prs)
    add_next_steps_slide(prs)
    add_takeaways_slide(prs)
    add_references_slide(prs)

    set_core_properties(prs)
    prs.save(PPTX_PATH)
    write_outline()
    validate_presentation(PPTX_PATH)
    print(PPTX_PATH)
    print(OUTLINE_PATH)


def set_core_properties(prs: Presentation) -> None:
    props = prs.core_properties
    props.title = "Automatic Pseudo-Mask Generation for Industrial Defect Synthesis"
    props.subject = "ImgGen v2 current progress: Qwen-guided morphology-aware auto-mask generation"
    props.author = "ImgGen v2 Research Project"
    props.keywords = "industrial anomaly, auto mask, pseudo-label, Qwen, PatchCore, FFT, diffusion"
    props.comments = "Generated from current v2 project artifacts on June 7, 2026."


def blank_slide(prs: Presentation, title: str, section: str | None = None) -> object:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = LIGHT
    add_rect(slide, 0, 0, 13.333, 0.12, NAVY, line=None)
    add_text(slide, 0.55, 0.27, 12.1, 0.48, title, 26, NAVY, bold=True)
    if section:
        add_text(slide, 10.55, 0.32, 2.15, 0.25, section.upper(), 9, TEAL_DARK, bold=True, align=PP_ALIGN.RIGHT)
    add_footer(slide, len(prs.slides))
    return slide


def add_footer(slide: object, number: int) -> None:
    add_line(slide, 0.55, 7.08, 12.75, 7.08, RGBColor(203, 213, 225), 0.7)
    add_text(slide, 0.58, 7.12, 4.5, 0.18, "ImgGen v2 | Auto-Mask Generation Progress", 8, MUTED)
    add_text(slide, 11.8, 7.12, 0.9, 0.18, f"{number:02d}", 8, MUTED, align=PP_ALIGN.RIGHT)


def add_title_slide(prs: Presentation, assets: dict[str, Path]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY
    add_rect(slide, 0, 0, 0.18, 7.5, TEAL, line=None)
    add_text(slide, 0.72, 0.68, 11.7, 0.55, "AUTOMATIC PSEUDO-MASK GENERATION", 13, RGBColor(94, 234, 212), bold=True)
    add_text(
        slide,
        0.72,
        1.28,
        7.45,
        1.55,
        "From normal + few defect images\nto morphology-aware masks",
        31,
        WHITE,
        bold=True,
    )
    add_text(
        slide,
        0.75,
        3.0,
        6.9,
        0.92,
        "No manual bounding-box or pixel-mask drawing in the user workflow",
        18,
        RGBColor(203, 213, 225),
    )
    add_pill(slide, 0.75, 4.22, 1.72, 0.38, "Qwen localization", TEAL)
    add_pill(slide, 2.58, 4.22, 1.68, 0.38, "16 candidates", PURPLE)
    add_pill(slide, 4.37, 4.22, 1.92, 0.38, "Uncertainty-aware", ORANGE)
    add_pill(slide, 6.4, 4.22, 1.55, 0.38, "Wood study", RED)
    add_text(slide, 0.75, 6.58, 6.9, 0.32, "Current research progress | June 2026", 12, RGBColor(148, 163, 184))
    add_picture_cover(slide, assets["title_visual"], 8.15, 0.68, 4.65, 5.95, radius=True)
    add_text(slide, 8.28, 6.78, 4.25, 0.25, "Current custom wood cases: scuff + vertical scratch + diagonal scratch", 9, RGBColor(148, 163, 184), align=PP_ALIGN.CENTER)


def add_manual_dependency_slide(prs: Presentation) -> None:
    slide = blank_slide(prs, "Why automatic masks matter", "Motivation")
    add_text(
        slide,
        0.7,
        0.92,
        11.9,
        0.58,
        "Common controllable industrial defect generation begins with a user-supplied bounding box or pixel mask.",
        21,
        NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        1.0,
        1.52,
        11.3,
        0.35,
        "This is not universal, but it is prevalent in spatially controlled and mask-conditioned synthesis.",
        11,
        MUTED,
        align=PP_ALIGN.CENTER,
    )

    draw_pipeline_box(slide, 0.75, 2.18, 2.35, 1.2, "Human operator", "Inspect defect image", ORANGE, icon="1")
    draw_pipeline_box(slide, 3.58, 2.18, 2.35, 1.2, "Draw location", "Bounding box or mask", ORANGE, icon="2")
    draw_pipeline_box(slide, 6.4, 2.18, 2.35, 1.2, "Generate defect", "Diffusion / inpainting", TEAL, icon="3")
    draw_pipeline_box(slide, 9.23, 2.18, 2.35, 1.2, "Train detector", "Image-mask pairs", PURPLE, icon="4")
    for x in (3.13, 5.95, 8.78):
        add_arrow(slide, x, 2.78, x + 0.38, 2.78, ORANGE if x < 5.95 else SLATE)

    add_callout(
        slide,
        0.85,
        4.05,
        3.55,
        1.36,
        "Annotation burden",
        "Pixel masks are expensive; rough boxes are cheaper but do not provide accurate defect morphology.",
        ORANGE,
    )
    add_callout(
        slide,
        4.88,
        4.05,
        3.55,
        1.36,
        "Alignment risk",
        "A supplied region can be misplaced, over-broad, or inconsistent with the generated anomaly.",
        RED,
    )
    add_callout(
        slide,
        8.91,
        4.05,
        3.55,
        1.36,
        "Deployment friction",
        "New parts and defect types require repeated manual spatial annotation before synthesis.",
        PURPLE,
    )
    add_text(
        slide,
        0.7,
        6.33,
        12.0,
        0.42,
        "Project response: infer the search region and purpose-specific pseudo-masks automatically from normal + defect images.",
        15,
        TEAL_DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_citation(
        slide,
        "Examples: Bounding Box-Guided Diffusion (Simoni & Pelosin, 2025); AnomalyDiffusion (Hu et al., AAAI 2024); "
        "DefFiller (Tai et al., 2024); SeaS (Dai et al., ICCV 2025).",
    )


def add_research_gap_slide(prs: Presentation) -> None:
    slide = blank_slide(prs, "Research gap and project objective", "Motivation")
    add_text(slide, 0.68, 0.92, 5.6, 0.36, "Existing controllable workflow", 18, ORANGE, bold=True)
    add_text(slide, 7.05, 0.92, 5.6, 0.36, "Target automatic workflow", 18, TEAL_DARK, bold=True)

    old = [
        ("Normal images", GREEN),
        ("Defect images", RED),
        ("Human box / mask", ORANGE),
        ("Generation model", PURPLE),
        ("Synthetic image-mask pairs", CYAN),
    ]
    new = [
        ("Normal images", GREEN),
        ("Few defect images", RED),
        ("Optional text", SLATE),
        ("Auto-mask generator", TEAL),
        ("Purpose-specific pseudo-masks", CYAN),
    ]
    draw_vertical_flow(slide, 0.85, 1.48, old, arrow_color=ORANGE)
    draw_vertical_flow(slide, 7.22, 1.48, new, arrow_color=TEAL)
    add_rect(slide, 6.48, 1.38, 0.03, 4.9, RGBColor(203, 213, 225), line=None)

    add_text(slide, 0.82, 6.38, 5.65, 0.42, "Spatial supervision must be drawn before generation.", 12, RED, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 7.16, 6.38, 5.65, 0.42, "Spatial supervision is inferred, scored, routed, and audited.", 12, TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)


def add_progress_snapshot_slide(prs: Presentation) -> None:
    slide = blank_slide(prs, "Current progress at a glance", "Status")
    cards = [
        ("3", "custom wood defects", "000 scuff, 001 vertical, 002 diagonal", RED),
        ("16", "candidate modes", "normal, texture, ridge, neural, fallback", PURPLE),
        ("8", "mask roles", "eval, train, uncertainty, generation", TEAL),
        ("3 / 3", "policy validation", "all custom wood manifest rows pass", GREEN),
    ]
    for i, (big, title, sub, color) in enumerate(cards):
        x = 0.65 + i * 3.08
        add_metric_card(slide, x, 1.05, 2.7, 1.42, big, title, sub, color)

    add_text(slide, 0.72, 2.9, 4.0, 0.35, "Implemented capabilities", 18, NAVY, bold=True)
    implemented = [
        "Qwen rough-region localization and sub-box fusion",
        "Clean-normal PatchCore-style candidate",
        "FFT texture suppression and structure-tensor ridges",
        "Morphology-aware hard/soft label routing",
        "Candidate disagreement uncertainty maps",
        "Mask-policy ablation and manifest validation",
    ]
    add_bullet_list(slide, 0.78, 3.35, 5.72, 2.7, implemented, 13, NAVY, bullet_color=TEAL)

    add_text(slide, 7.0, 2.9, 4.0, 0.35, "Current evidence", 18, NAVY, bold=True)
    evidence = [
        ("000", "patchcore_guided", "multi_scuff", "training_soft", ORANGE),
        ("001", "ensemble_consensus", "scratch_band", "training_medium", GREEN),
        ("002", "ensemble_consensus", "scratch_band", "training_medium", GREEN),
    ]
    y = 3.36
    for sample, candidate, morphology, mask, color in evidence:
        add_result_row(slide, 7.02, y, 5.55, 0.67, sample, candidate, morphology, mask, color)
        y += 0.78
    add_callout(slide, 7.02, 5.84, 5.55, 0.75, "Research status", "Architecture implemented; broader quantitative validation is still in progress.", PURPLE)


def add_architecture_slide(prs: Presentation) -> None:
    slide = blank_slide(prs, "Auto-mask generation architecture", "System")
    stages = [
        ("Inputs", "normal + defect\n+ optional text", GREEN),
        ("Qwen", "rough region\nand morphology cues", CYAN),
        ("Candidates", "16 mask / heatmap\nhypotheses", PURPLE),
        ("Policy", "score + ensemble\n+ morphology route", ORANGE),
        ("Outputs", "eval / train /\nuncertainty / inpaint", TEAL),
    ]
    x_positions = [0.55, 3.05, 5.55, 8.05, 10.55]
    for i, ((title, sub, color), x) in enumerate(zip(stages, x_positions)):
        draw_stage_card(slide, x, 1.15, 2.22, 1.48, title, sub, color, i + 1)
        if i < len(stages) - 1:
            add_arrow(slide, x + 2.23, 1.9, x + 2.48, 1.9, SLATE)

    add_text(slide, 0.72, 3.05, 12.0, 0.38, "Two complementary reasoning paths", 18, NAVY, bold=True)
    add_rect(slide, 0.75, 3.55, 5.7, 2.36, WHITE, line=RGBColor(203, 213, 225), radius=True)
    add_text(slide, 1.0, 3.82, 5.2, 0.32, "A. Spatial / semantic path", 16, CYAN, bold=True)
    add_bullet_list(
        slide,
        1.0,
        4.25,
        5.05,
        1.32,
        [
            "Qwen identifies a padded search fence, not a final mask",
            "Description cues suggest multi_scuff vs scratch_band",
            "Region constraints suppress unrelated background",
        ],
        12,
        NAVY,
        bullet_color=CYAN,
    )
    add_rect(slide, 6.85, 3.55, 5.7, 2.36, WHITE, line=RGBColor(203, 213, 225), radius=True)
    add_text(slide, 7.1, 3.82, 5.2, 0.32, "B. Pixel / texture evidence path", 16, PURPLE, bold=True)
    add_bullet_list(
        slide,
        7.1,
        4.25,
        5.05,
        1.32,
        [
            "Compare defect image against clean-normal texture memory",
            "Measure residual, frequency disruption, and ridge geometry",
            "Fuse candidate support and expose disagreement",
        ],
        12,
        NAVY,
        bullet_color=PURPLE,
    )
    add_text(slide, 0.85, 6.35, 11.7, 0.38, "Design principle: use the VLM to narrow the search; use image evidence to decide pixels.", 14, TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)


def add_input_contract_slide(prs: Presentation, assets: dict[str, Path]) -> None:
    slide = blank_slide(prs, "User input contract", "System")
    add_text(slide, 0.72, 0.9, 12.0, 0.42, "The normal workflow remains automatic and scarce-data friendly.", 19, NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_picture_cover(slide, assets["normal"], 0.72, 1.58, 3.2, 3.55, radius=True)
    add_picture_cover(slide, assets["defect_000"], 5.06, 1.58, 3.2, 3.55, radius=True)
    add_rect(slide, 9.4, 1.58, 3.2, 3.55, WHITE, line=RGBColor(203, 213, 225), radius=True)
    add_text(slide, 9.72, 2.03, 2.55, 0.35, "Optional description", 16, SLATE, bold=True, align=PP_ALIGN.CENTER)
    add_text(
        slide,
        9.76,
        2.7,
        2.47,
        1.2,
        "“a rectangular scuffed patch with many faint short scratches and rubbed light damage”",
        15,
        NAVY,
        italic=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_pill(slide, 1.28, 5.35, 2.08, 0.4, "Clean normal images", GREEN)
    add_pill(slide, 5.61, 5.35, 2.08, 0.4, "Few defect images", RED)
    add_pill(slide, 9.95, 5.35, 2.08, 0.4, "Optional text", SLATE)
    add_text(slide, 0.9, 6.18, 11.6, 0.44, "Not required: manual bounding boxes, manual segmentation masks, or per-image clicks.", 17, TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)


def add_qwen_localization_slide(prs: Presentation, assets: dict[str, Path], records: dict[str, dict]) -> None:
    slide = blank_slide(prs, "Qwen localization: a search fence, not a mask oracle", "Stage 1")
    add_text(slide, 0.72, 0.9, 12.0, 0.35, "Qwen returns a rough padded region. Pixel-level candidates are computed only inside that region.", 15, NAVY, align=PP_ALIGN.CENTER)
    samples = ["000", "001", "002"]
    for i, sample in enumerate(samples):
        x = 0.7 + i * 4.22
        add_picture_cover(slide, assets[f"bbox_{sample}"], x, 1.48, 3.65, 3.55, radius=True)
        region = records[sample]["region_xyxy"]
        add_text(slide, x, 5.18, 3.65, 0.3, f"{sample}: Qwen region {region}", 10, SLATE, bold=True, align=PP_ALIGN.CENTER)
        morph = records[sample]["settings"]["quality_morphology"]
        color = ORANGE if morph == "multi_scuff" else GREEN
        add_pill(slide, x + 0.88, 5.58, 1.9, 0.36, morph, color)
    add_callout(
        slide,
        0.85,
        6.18,
        11.65,
        0.56,
        "Why this matters",
        "A rough semantic region is easier to obtain automatically; the candidate bank then handles fine boundaries and textured surfaces.",
        CYAN,
    )


def add_candidate_bank_slide(prs: Presentation) -> None:
    slide = blank_slide(prs, "Candidate bank: 16 complementary hypotheses", "Stage 1")
    groups = [
        (
            "Normal-reference evidence",
            ["normal_anomaly", "nearest_normal_residual", "patchcore_guided", "normal_residual_fusion"],
            GREEN,
        ),
        (
            "Texture / frequency evidence",
            ["fft_texture_suppression", "soft_patch", "scuff_cluster", "multi_scuff_fusion"],
            ORANGE,
        ),
        (
            "Geometry / ridge evidence",
            ["scratch_band_clean", "structure_tensor_ridge", "multi_linear", "residual"],
            PURPLE,
        ),
        (
            "Neural + robust fallbacks",
            ["sam2_heatmap", "support_constrained_fusion", "pixel", "procedural"],
            CYAN,
        ),
    ]
    positions = [(0.65, 1.08), (6.75, 1.08), (0.65, 3.83), (6.75, 3.83)]
    for (title, items, color), (x, y) in zip(groups, positions):
        add_rect(slide, x, y, 5.9, 2.28, WHITE, line=RGBColor(203, 213, 225), radius=True)
        add_rect(slide, x, y, 0.12, 2.28, color, line=None)
        add_text(slide, x + 0.3, y + 0.22, 5.25, 0.32, title, 16, color, bold=True)
        yy = y + 0.72
        for j, item in enumerate(items):
            col = j % 2
            row = j // 2
            add_pill(slide, x + 0.34 + col * 2.68, yy + row * 0.57, 2.4, 0.38, item, color, font_size=9)
    add_text(
        slide,
        0.85,
        6.43,
        11.7,
        0.34,
        "No single candidate is trusted universally. Selection depends on image evidence, morphology, area, leakage, and candidate agreement.",
        13,
        NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )


def add_candidate_scoring_slide(prs: Presentation) -> None:
    slide = blank_slide(prs, "Candidate scoring and selection", "Stage 1")
    add_text(slide, 0.72, 0.9, 12.0, 0.34, "Each candidate is treated as a hypothesis with diagnostics, not as ground truth.", 16, NAVY, bold=True, align=PP_ALIGN.CENTER)

    factors = [
        ("Residual support", "Does the candidate overlap high anomaly evidence?", GREEN),
        ("Area sanity", "Is it too small, over-broad, or nearly rectangular?", ORANGE),
        ("Border leakage", "Does it spill outside the Qwen search fence?", RED),
        ("Morphology fit", "Does shape match scuff, scratch-band, or stroke?", PURPLE),
        ("Agreement", "Do independent candidate families support it?", CYAN),
        ("Warnings", "Low contrast, fragmentation, or texture-alignment risk", SLATE),
    ]
    for i, (title, desc, color) in enumerate(factors):
        x = 0.7 + (i % 3) * 4.18
        y = 1.45 + (i // 3) * 1.62
        add_factor_card(slide, x, y, 3.7, 1.27, title, desc, color)

    add_rect(slide, 1.18, 5.05, 11.0, 1.15, NAVY, line=None, radius=True)
    add_text(slide, 1.55, 5.27, 2.1, 0.3, "Selection logic", 15, RGBColor(94, 234, 212), bold=True)
    add_text(slide, 3.42, 5.22, 8.35, 0.5, "score candidates → optional ensemble consensus → morphology policy → purpose-specific masks", 17, WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 1.4, 6.42, 10.5, 0.32, "Warning score caps prevent risky candidates from appearing as perfect 1.0 results.", 12, RED, bold=True, align=PP_ALIGN.CENTER)


def add_scuff_case_slide(prs: Presentation, assets: dict[str, Path], records: dict[str, dict]) -> None:
    slide = blank_slide(prs, "Case 000: broad, low-contrast multi-scuff", "Wood Study")
    add_picture_cover(slide, assets["defect_000"], 0.62, 1.12, 2.45, 2.62, radius=True)
    add_text(slide, 0.72, 3.9, 2.25, 0.42, "Problem", 15, RED, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 0.72, 4.3, 2.25, 1.25, "Texture disruption is broad and faint.\nA single crisp line is technically dishonest.", 12, NAVY, align=PP_ALIGN.CENTER)

    heatmaps = [
        ("PatchCore", assets["000_patchcore"]),
        ("Nearest residual", assets["000_nearest"]),
        ("FFT suppression", assets["000_fft"]),
        ("Normal anomaly", assets["000_normal"]),
    ]
    for i, (label, path) in enumerate(heatmaps):
        x = 3.35 + (i % 2) * 2.45
        y = 1.12 + (i // 2) * 2.25
        add_picture_cover(slide, path, x, y, 2.05, 1.65, radius=True)
        add_text(slide, x, y + 1.7, 2.05, 0.25, label, 10, SLATE, bold=True, align=PP_ALIGN.CENTER)

    add_arrow(slide, 8.12, 3.0, 8.65, 3.0, SLATE)
    add_picture_cover(slide, assets["000_training_soft_overlay"], 8.82, 1.12, 3.65, 3.72, radius=True)
    add_text(slide, 8.88, 4.96, 3.5, 0.35, "Calibrated training_soft", 15, TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)
    add_pill(slide, 9.4, 5.47, 2.55, 0.38, "soft_mask_only", ORANGE)
    add_text(slide, 3.35, 5.7, 4.75, 0.55, "Selected candidate: patchcore_guided\nTraining target: fused calibrated soft evidence", 13, NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 0.82, 6.48, 11.7, 0.3, "Current policy: hard mask only for compact pseudo-evaluation; soft support for training and generation.", 13, TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)


def add_scratch_cases_slide(prs: Presentation, assets: dict[str, Path], records: dict[str, dict]) -> None:
    slide = blank_slide(prs, "Cases 001 and 002: clear line-like scratches", "Wood Study")
    add_text(slide, 0.72, 0.9, 12.0, 0.32, "Line geometry is stable enough for hard pseudo-label training.", 16, NAVY, bold=True, align=PP_ALIGN.CENTER)
    samples = [
        ("001", "Vertical scratch", assets["defect_001"], assets["001_ridge"], assets["001_training_medium_overlay"]),
        ("002", "Diagonal scratch band", assets["defect_002"], assets["002_ridge"], assets["002_training_medium_overlay"]),
    ]
    for row, (sample, title, defect, ridge, output) in enumerate(samples):
        y = 1.45 + row * 2.65
        add_text(slide, 0.65, y + 0.1, 1.0, 0.35, sample, 22, PURPLE, bold=True, align=PP_ALIGN.CENTER)
        add_picture_cover(slide, defect, 1.55, y, 2.25, 2.15, radius=True)
        add_picture_cover(slide, ridge, 4.25, y, 2.25, 2.15, radius=True)
        add_arrow(slide, 6.67, y + 1.08, 7.1, y + 1.08, SLATE)
        add_picture_cover(slide, output, 7.3, y, 2.65, 2.15, radius=True)
        add_rect(slide, 10.25, y, 2.42, 2.15, WHITE, line=RGBColor(203, 213, 225), radius=True)
        add_text(slide, 10.47, y + 0.24, 1.98, 0.32, title, 14, NAVY, bold=True, align=PP_ALIGN.CENTER)
        add_pill(slide, 10.65, y + 0.82, 1.62, 0.34, "hard_mask_ok", GREEN, font_size=9)
        add_pill(slide, 10.48, y + 1.3, 1.95, 0.34, "training_medium", TEAL, font_size=9)
    add_text(slide, 1.78, 6.82, 1.8, 0.22, "Defect image", 9, MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, 4.48, 6.82, 1.8, 0.22, "Ridge evidence", 9, MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, 7.7, 6.82, 1.8, 0.22, "Selected training mask", 9, MUTED, align=PP_ALIGN.CENTER)


def add_policy_routing_slide(prs: Presentation) -> None:
    slide = blank_slide(prs, "Morphology-aware label policy", "Stage 1")
    add_text(slide, 0.72, 0.9, 12.0, 0.34, "The same binary mask policy is not appropriate for every defect.", 17, NAVY, bold=True, align=PP_ALIGN.CENTER)

    columns = ["Morphology", "Evidence pattern", "Label policy", "Training mask", "Generation mask"]
    rows = [
        ["multi_scuff", "broad + faint + uncertain boundary", "soft_mask_only", "training_soft", "inpaint_soft"],
        ["scratch_band", "elongated ridge / line support", "hard_mask_ok", "training_medium", "inpaint_soft"],
        ["single_stroke", "one narrow continuous trace", "hard_mask_ok", "training_medium", "inpaint_soft"],
        ["crack_band", "thin branching or ridge-like trace", "hard_mask_ok", "training_medium", "inpaint_soft"],
    ]
    add_table(slide, 0.65, 1.48, 12.05, 3.45, columns, rows, [1.65, 3.15, 1.75, 1.75, 1.75])
    add_callout(slide, 0.9, 5.3, 5.45, 1.05, "Hard-mask route", "Use when geometry is line-like and candidate agreement is stable.", GREEN)
    add_callout(slide, 6.95, 5.3, 5.45, 1.05, "Soft-mask route", "Use when damage changes texture gradually and boundaries are ambiguous.", ORANGE)
    add_text(slide, 1.0, 6.62, 11.4, 0.3, "Automatic masks are explicitly labeled as pseudo-labels, not human ground truth.", 12, RED, bold=True, align=PP_ALIGN.CENTER)


def add_mask_variants_slide(prs: Presentation, assets: dict[str, Path]) -> None:
    slide = blank_slide(prs, "Purpose-specific mask outputs", "Stage 1")
    add_text(slide, 0.72, 0.9, 12.0, 0.32, "One selected defect produces multiple masks because evaluation, training, and diffusion need different support.", 15, NAVY, align=PP_ALIGN.CENTER)
    variants = [
        ("eval_tight", "Compact pseudo-evaluation", assets["000_eval_overlay"], RED),
        ("training_medium", "Hard line training support", assets["001_training_medium_overlay"], GREEN),
        ("training_soft", "Soft fuzzy scuff target", assets["000_training_soft_overlay"], ORANGE),
        ("uncertainty_map", "Candidate disagreement", assets["000_uncertainty_overlay"], PURPLE),
        ("positive_core", "High-confidence positive", assets["000_core_overlay"], TEAL),
        ("possible_region", "Plausible support envelope", assets["000_possible_overlay"], CYAN),
        ("inpaint_soft", "Diffusion blending support", assets["000_inpaint_overlay"], SLATE),
    ]
    for i, (name, desc, path, color) in enumerate(variants):
        x = 0.48 + i * 1.82
        add_picture_cover(slide, path, x, 1.5, 1.58, 2.55, radius=True)
        add_text(slide, x, 4.18, 1.58, 0.28, name, 10, color, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, x + 0.03, 4.52, 1.52, 0.67, desc, 9, NAVY, align=PP_ALIGN.CENTER)
    add_rect(slide, 0.75, 5.62, 11.85, 0.85, NAVY, line=None, radius=True)
    add_text(slide, 1.05, 5.88, 11.2, 0.34, "Separation of roles prevents a broad inpainting mask from becoming a fake crisp training label.", 14, WHITE, bold=True, align=PP_ALIGN.CENTER)


def add_uncertainty_slide(prs: Presentation, assets: dict[str, Path]) -> None:
    slide = blank_slide(prs, "Uncertainty-aware pseudo-label learning", "Stage 1")
    add_picture_cover(slide, assets["000_training_soft_overlay"], 0.7, 1.3, 3.25, 3.8, radius=True)
    add_picture_cover(slide, assets["000_uncertainty_overlay"], 4.35, 1.3, 3.25, 3.8, radius=True)
    add_rect(slide, 8.0, 1.3, 4.55, 3.8, WHITE, line=RGBColor(203, 213, 225), radius=True)
    add_text(slide, 8.28, 1.68, 4.0, 0.35, "Loss weighting", 18, PURPLE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 8.38, 2.38, 3.78, 0.75, "w(x) = 1 − U(x) · (1 − α)", 24, NAVY, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, 8.42, 3.25, 3.7, 0.52, "U(x): uncertainty map\nα = 0.25 in Phase 3 and Phase 5", 14, SLATE, align=PP_ALIGN.CENTER)
    add_pill(slide, 8.95, 4.22, 2.65, 0.42, "uncertainty_loss_weight = 0.25", PURPLE, font_size=9)
    add_text(slide, 1.15, 5.3, 2.35, 0.3, "training_soft", 13, ORANGE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 4.82, 5.3, 2.35, 0.3, "uncertainty_map", 13, PURPLE, bold=True, align=PP_ALIGN.CENTER)
    add_bullet_list(
        slide,
        0.95,
        5.78,
        11.45,
        0.88,
        [
            "High-confidence core and background keep full supervision.",
            "Candidate-disagreement regions contribute less, reducing damage from uncertain automatic boundaries.",
        ],
        12,
        NAVY,
        bullet_color=PURPLE,
    )


def add_wood_results_slide(prs: Presentation, assets: dict[str, Path], records: dict[str, dict]) -> None:
    slide = blank_slide(prs, "Current wood auto-mask result", "Results")
    add_picture_contain(slide, assets["wood_visual"], 0.45, 1.05, 8.15, 5.95)
    add_rect(slide, 8.85, 1.05, 3.95, 5.8, WHITE, line=RGBColor(203, 213, 225), radius=True)
    add_text(slide, 9.15, 1.35, 3.35, 0.35, "Per-image routing", 18, NAVY, bold=True, align=PP_ALIGN.CENTER)
    result_rows = [
        ("000", "patchcore_guided", "multi_scuff", "training_soft", ORANGE),
        ("001", "ensemble_consensus", "scratch_band", "training_medium", GREEN),
        ("002", "ensemble_consensus", "scratch_band", "training_medium", GREEN),
    ]
    y = 1.98
    for sample, candidate, morph, mask, color in result_rows:
        add_result_row(slide, 9.12, y, 3.4, 0.86, sample, candidate, morph, mask, color, compact=True)
        y += 1.02
    add_text(slide, 9.15, 5.2, 3.35, 0.3, "Mask pixels", 14, SLATE, bold=True, align=PP_ALIGN.CENTER)
    pixel_lines = [
        "000: eval 6,737 | soft 21,498",
        "001: eval 14,349 | medium 19,576",
        "002: eval 12,252 | medium 36,039",
    ]
    add_bullet_list(slide, 9.18, 5.55, 3.3, 0.96, pixel_lines, 10, NAVY, bullet_color=TEAL)


def add_candidate_comparison_slide(prs: Presentation, assets: dict[str, Path]) -> None:
    slide = blank_slide(prs, "Visual comparison across candidate methods", "Results")
    add_picture_contain(slide, assets["candidate_comparison"], 0.35, 1.03, 12.65, 4.95)
    add_callout(slide, 0.8, 6.05, 3.75, 0.62, "000 multi-scuff", "PatchCore-guided support avoids forcing one artificial line.", ORANGE)
    add_callout(slide, 4.8, 6.05, 3.75, 0.62, "001 vertical", "Ensemble consensus preserves a narrow vertical trace.", GREEN)
    add_callout(slide, 8.8, 6.05, 3.75, 0.62, "002 diagonal", "Ridge and residual methods agree on the diagonal band.", GREEN)


def add_ablation_slide(prs: Presentation, assets: dict[str, Path], rows: list[dict[str, str]]) -> None:
    slide = blank_slide(prs, "Mask-policy ablation", "Evaluation")
    add_text(slide, 0.72, 0.9, 12.0, 0.32, "Diagnostic critic comparison: hard binary vs candidate-vote soft vs calibrated soft.", 15, NAVY, align=PP_ALIGN.CENTER)
    chart_data = ChartData()
    chart_data.categories = ["000 scuff", "001 vertical", "002 diagonal"]
    policies = ["hard_binary", "vote_soft", "calibrated_soft"]
    for policy in policies:
        values = [float(next(r["critic_score"] for r in rows if r["image"] == f"{sample}.png" and r["policy"] == policy)) for sample in ("000", "001", "002")]
        chart_data.add_series(policy, values)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.7), Inches(1.45), Inches(7.3), Inches(4.5), chart_data).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.value_axis.minimum_scale = 0.45
    chart.value_axis.maximum_scale = 0.76
    chart.value_axis.major_unit = 0.05
    chart.value_axis.has_major_gridlines = True
    chart.category_axis.tick_labels.font.size = Pt(11)
    chart.value_axis.tick_labels.font.size = Pt(9)
    colors = [GREEN, ORANGE, PURPLE]
    for series, color in zip(chart.series, colors):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = color
        series.format.line.color.rgb = color

    add_rect(slide, 8.35, 1.45, 4.25, 4.5, WHITE, line=RGBColor(203, 213, 225), radius=True)
    add_text(slide, 8.65, 1.75, 3.65, 0.35, "Per-image winner", 17, NAVY, bold=True, align=PP_ALIGN.CENTER)
    winners = [
        ("000", "vote_soft", "0.6703", ORANGE),
        ("001", "hard_binary", "0.6955", GREEN),
        ("002", "hard_binary", "0.7224", GREEN),
    ]
    y = 2.4
    for sample, policy, score, color in winners:
        add_rect(slide, 8.7, y, 3.55, 0.76, PALE, line=RGBColor(226, 232, 240), radius=True)
        add_text(slide, 8.92, y + 0.16, 0.52, 0.28, sample, 16, color, bold=True)
        add_text(slide, 9.48, y + 0.12, 1.65, 0.32, policy, 12, NAVY, bold=True)
        add_text(slide, 11.2, y + 0.12, 0.82, 0.32, score, 12, SLATE, bold=True, align=PP_ALIGN.RIGHT)
        y += 0.93
    add_text(slide, 8.72, 5.35, 3.5, 0.43, "Interpretation", 14, PURPLE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 8.72, 5.72, 3.5, 0.88, "Soft support is justified for fuzzy scuffs; clear scratches still prefer hard masks.", 12, NAVY, align=PP_ALIGN.CENTER)
    add_citation(slide, "Critic is diagnostic pseudo-label evidence, not human ground truth.")


def add_manifest_validation_slide(prs: Presentation) -> None:
    slide = blank_slide(prs, "Mask-policy validation in the Phase 5 manifest", "Evaluation")
    add_metric_card(slide, 0.75, 1.1, 2.75, 1.42, "3 / 3", "rows pass", "custom wood manifest", GREEN)
    add_metric_card(slide, 3.85, 1.1, 2.75, 1.42, "1", "soft scuff row", "multi_scuff → training_soft", ORANGE)
    add_metric_card(slide, 6.95, 1.1, 2.75, 1.42, "2", "hard scratch rows", "scratch_band → training_medium", TEAL)
    add_metric_card(slide, 10.05, 1.1, 2.55, 1.42, "0", "policy mismatches", "current custom set", PURPLE)

    columns = ["Image", "Split", "Morphology", "Policy", "Observed training role", "Status"]
    rows = [
        ["002.png", "adaptation", "scratch_band", "hard_mask_ok", "training_medium", "PASS"],
        ["000.png", "adaptation", "multi_scuff", "soft_mask_only", "training_soft", "PASS"],
        ["001.png", "held_out", "scratch_band", "hard_mask_ok", "training_medium", "PASS"],
    ]
    add_table(slide, 0.68, 2.92, 12.0, 2.62, columns, rows, [1.2, 1.35, 1.65, 1.75, 2.35, 0.85])
    add_text(slide, 0.95, 5.9, 11.5, 0.52, "This confirms that the routing decision survives into the actual training/evaluation manifest.", 15, TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 0.95, 6.47, 11.5, 0.28, "Artifact: reports/auto_mask_phase9_wood/phase5/qwen/mask_policy_validation.md", 9, MUTED, align=PP_ALIGN.CENTER)


def add_downstream_integration_slide(prs: Presentation, assets: dict[str, Path]) -> None:
    slide = blank_slide(prs, "How auto-masks connect to generation and evaluation", "Integration")
    stages = [
        ("Auto masks", "eval_tight\ntraining_soft / medium", TEAL),
        ("Phase 3", "adapter loss with\nuncertainty weighting", PURPLE),
        ("Phase 4", "SD1.5 generation\nwith inpaint_soft", CYAN),
        ("Phase 11", "TF-IDG-lite critic\ncoverage + texture", ORANGE),
        ("Phase 5", "selector + U-Net /\nPatchCore evaluation", GREEN),
    ]
    for i, (title, sub, color) in enumerate(stages):
        x = 0.48 + i * 2.56
        draw_stage_card(slide, x, 1.15, 2.28, 1.35, title, sub, color, i + 1)
        if i < 4:
            add_arrow(slide, x + 2.29, 1.84, x + 2.51, 1.84, SLATE)

    add_picture_cover(slide, assets["wood_end_to_end"], 0.72, 2.95, 7.0, 3.5, radius=True)
    add_rect(slide, 8.05, 2.95, 4.55, 3.5, WHITE, line=RGBColor(203, 213, 225), radius=True)
    add_text(slide, 8.35, 3.25, 3.95, 0.35, "Current critic diagnosis", 17, NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_metric_card(slide, 8.4, 3.85, 1.75, 1.1, "540", "wood samples", "scored", CYAN)
    add_metric_card(slide, 10.4, 3.85, 1.75, 1.1, "74", "accepted", "current thresholds", GREEN)
    add_text(slide, 8.48, 5.2, 3.78, 0.62, "466 rejected mainly for\nlow mask coverage", 18, RED, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 8.38, 5.92, 3.95, 0.33, "Next: regenerate or reject coverage < 0.35", 11, ORANGE, bold=True, align=PP_ALIGN.CENTER)
    add_citation(slide, "TF-IDG inspiration: Xu et al., Training-Free Industrial Defect Generation with Diffusion Models, ICCV 2025.")


def add_strengths_limitations_slide(prs: Presentation) -> None:
    slide = blank_slide(prs, "Professional assessment: advantages and limitations", "Review")
    add_rect(slide, 0.65, 1.1, 5.95, 5.65, WHITE, line=RGBColor(187, 247, 208), radius=True)
    add_rect(slide, 6.75, 1.1, 5.95, 5.65, WHITE, line=RGBColor(254, 202, 202), radius=True)
    add_text(slide, 0.95, 1.43, 5.35, 0.4, "Advantages", 21, GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 7.05, 1.43, 5.35, 0.4, "Limitations", 21, RED, bold=True, align=PP_ALIGN.CENTER)
    strengths = [
        "No manual box or mask drawing in the normal user workflow",
        "Normal-reference evidence reduces texture-only false positives",
        "Morphology routing avoids one-size-fits-all binary labels",
        "Purpose-specific masks separate evaluation, training, and inpainting",
        "Uncertainty is exposed and used in the loss",
        "Reports preserve candidate diagnostics and reproducibility",
    ]
    limitations = [
        "Automatic masks remain pseudo-labels, not human annotations",
        "Broad low-contrast scuffs still have ambiguous boundaries",
        "Candidate thresholds are currently engineering-calibrated",
        "Evidence is strongest on only three custom wood cases",
        "Qwen localization can fail on unseen materials or descriptions",
        "Final claims need human-mask subset evaluation and stronger evaluators",
    ]
    add_bullet_list(slide, 0.95, 2.1, 5.2, 4.18, strengths, 12, NAVY, bullet_color=GREEN, spacing=7)
    add_bullet_list(slide, 7.05, 2.1, 5.2, 4.18, limitations, 12, NAVY, bullet_color=RED, spacing=7)


def add_next_steps_slide(prs: Presentation) -> None:
    slide = blank_slide(prs, "Next research upgrades", "Roadmap")
    steps = [
        ("1", "Close the generation loop", "Regenerate or reject samples with TF-IDG mask coverage < 0.35.", ORANGE),
        ("2", "Human-labeled audit subset", "Label a small set to measure mask IoU, boundary F-score, and calibration.", RED),
        ("3", "Learned candidate calibration", "Fit thresholds or a lightweight ranker across materials without leakage.", PURPLE),
        ("4", "Broader material validation", "Run metal, tile, fabric, and plastic with morphology-specific reporting.", CYAN),
        ("5", "Stronger downstream evidence", "PatchCore-ResNet, DRAEM-style reconstruction, repeated supervised seeds.", GREEN),
    ]
    y = 1.08
    for number, title, desc, color in steps:
        add_rect(slide, 0.72, y, 11.9, 0.94, WHITE, line=RGBColor(203, 213, 225), radius=True)
        add_pill(slide, 0.94, y + 0.24, 0.58, 0.42, number, color, font_size=12)
        add_text(slide, 1.75, y + 0.18, 3.3, 0.32, title, 15, color, bold=True)
        add_text(slide, 5.05, y + 0.17, 7.15, 0.5, desc, 12, NAVY)
        y += 1.1
    add_text(slide, 0.9, 6.7, 11.5, 0.28, "Priority: validate reliability before adding more candidate modes.", 13, TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)


def add_takeaways_slide(prs: Presentation) -> None:
    slide = blank_slide(prs, "Key takeaways", "Conclusion")
    takeaways = [
        ("Automatic spatial supervision", "The workflow replaces manual box/mask drawing with Qwen-guided candidate search.", TEAL),
        ("Morphology-aware labels", "Fuzzy scuffs receive soft targets; line scratches retain hard masks.", ORANGE),
        ("Honest uncertainty", "Candidate disagreement becomes an uncertainty map and lower loss weight.", PURPLE),
        ("Traceable evidence", "Every selection has candidate scores, mask variants, reports, and manifest checks.", GREEN),
        ("Next bottleneck", "Generation occupancy and stronger validation now matter more than adding heuristics.", RED),
    ]
    y = 1.02
    for i, (title, desc, color) in enumerate(takeaways):
        x = 0.75 if i < 3 else 6.75
        yy = y + (i if i < 3 else i - 3) * 1.7
        width = 5.85
        add_rect(slide, x, yy, width, 1.28, WHITE, line=RGBColor(203, 213, 225), radius=True)
        add_rect(slide, x, yy, 0.12, 1.28, color, line=None)
        add_text(slide, x + 0.35, yy + 0.2, width - 0.62, 0.32, title, 16, color, bold=True)
        add_text(slide, x + 0.35, yy + 0.63, width - 0.62, 0.43, desc, 11, NAVY)
    add_rect(slide, 6.75, 4.72, 5.85, 1.82, NAVY, line=None, radius=True)
    add_text(slide, 7.05, 5.03, 5.25, 0.4, "Current research claim", 17, RGBColor(94, 234, 212), bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 7.08, 5.5, 5.18, 0.68, "A morphology-aware, uncertainty-aware automatic pseudo-mask pipeline is implemented and visually validated on custom wood defects.", 13, WHITE, align=PP_ALIGN.CENTER)


def add_references_slide(prs: Presentation) -> None:
    slide = blank_slide(prs, "Selected references", "References")
    refs = [
        ("Simoni & Pelosin (2025)", "Bounding Box-Guided Diffusion for Synthesizing Industrial Images and Segmentation Map.", "https://arxiv.org/abs/2505.03623"),
        ("Hu et al. (AAAI 2024)", "AnomalyDiffusion: Few-Shot Anomaly Image Generation with Diffusion Model.", "https://arxiv.org/abs/2312.05767"),
        ("Tai et al. (2024)", "DefFiller: Mask-Conditioned Diffusion for Salient Steel Surface Defect Generation.", "https://arxiv.org/abs/2412.15570"),
        ("Dai et al. (ICCV 2025)", "SeaS: Few-shot Industrial Anomaly Image Generation with Separation and Sharing Fine-tuning.", "https://openaccess.thecvf.com/content/ICCV2025/html/Dai_SeaS_Few-shot_Industrial_Anomaly_Image_Generation_with_Separation_and_Sharing_ICCV_2025_paper.html"),
        ("Xu et al. (ICCV 2025)", "Training-Free Industrial Defect Generation with Diffusion Models.", "https://openaccess.thecvf.com/content/ICCV2025/papers/Xu_Training-Free_Industrial_Defect_Generation_with_Diffusion_Models_ICCV_2025_paper.pdf"),
        ("Jin et al. (CVPR 2025)", "Dual-Interrelated Diffusion Model for Few-Shot Anomaly Image Generation.", "https://openaccess.thecvf.com/content/CVPR2025/html/Jin_Dual-Interrelated_Diffusion_Model_for_Few-Shot_Anomaly_Image_Generation_CVPR_2025_paper.html"),
    ]
    y = 1.02
    for author, title, url in refs:
        add_text(slide, 0.82, y, 2.35, 0.3, author, 11, TEAL_DARK, bold=True)
        add_text(slide, 3.0, y, 9.45, 0.3, title, 11, NAVY)
        add_text(slide, 3.0, y + 0.32, 9.45, 0.24, url, 8, MUTED)
        y += 0.86
    add_rect(slide, 0.82, 6.28, 11.65, 0.5, PALE, line=RGBColor(226, 232, 240), radius=True)
    add_text(slide, 1.02, 6.42, 11.25, 0.2, "Project artifacts and numerical results are generated from the local ImgGen v2 workspace.", 9, SLATE, align=PP_ALIGN.CENTER)


def build_assets() -> dict[str, Path]:
    assets: dict[str, Path] = {
        "normal": DATA_ROOT / "train" / "good" / "000.png",
        "defect_000": DATA_ROOT / "test" / "scratch" / "000.png",
        "defect_001": DATA_ROOT / "test" / "scratch" / "001.png",
        "defect_002": DATA_ROOT / "test" / "scratch" / "002.png",
        "wood_visual": REPORT_ROOT / "auto_masks" / "qwen" / "phase9_wood_visual_result.png",
        "candidate_comparison": REPORT_ROOT / "auto_masks" / "qwen" / "contact_sheet_candidate_comparison.png",
        "wood_end_to_end": TARGETED_REPORT / "wood_visual_result" / "wood_end_to_end_visual_result.png",
    }
    records = load_auto_mask_records()
    for sample in ("000", "001", "002"):
        bbox_path = ASSET_DIR / f"{sample}_qwen_bbox.png"
        draw_bbox_asset(
            Path(records[sample]["image_path"]),
            records[sample]["region_xyxy"],
            bbox_path,
            f"Qwen search region: {sample}",
        )
        assets[f"bbox_{sample}"] = bbox_path

    title_visual = ASSET_DIR / "title_visual.png"
    make_title_triptych(title_visual)
    assets["title_visual"] = title_visual

    masks = MASK_ROOT / "mask_variants" / "custom_part" / "scratch"
    assets.update(
        {
            "000_eval_overlay": masks / "custom_part_scratch_000_eval_tight_overlay.png",
            "000_training_soft_overlay": masks / "custom_part_scratch_000_training_soft_overlay.png",
            "000_uncertainty_overlay": masks / "custom_part_scratch_000_uncertainty_map_overlay.png",
            "000_core_overlay": masks / "custom_part_scratch_000_positive_core_overlay.png",
            "000_possible_overlay": masks / "custom_part_scratch_000_possible_region_overlay.png",
            "000_inpaint_overlay": masks / "custom_part_scratch_000_inpaint_soft_overlay.png",
            "001_training_medium_overlay": masks / "custom_part_scratch_001_training_medium_overlay.png",
            "002_training_medium_overlay": masks / "custom_part_scratch_002_training_medium_overlay.png",
        }
    )
    candidate = MASK_ROOT / "masks" / "custom_part" / "scratch"
    raw_heatmaps = {
        "000_patchcore": candidate / "custom_part_scratch_000_patchcore_guided_heatmap.png",
        "000_nearest": candidate / "custom_part_scratch_000_nearest_normal_residual_heatmap.png",
        "000_fft": candidate / "custom_part_scratch_000_fft_texture_suppression_heatmap.png",
        "000_normal": candidate / "custom_part_scratch_000_normal_anomaly_heatmap.png",
        "001_ridge": candidate / "custom_part_scratch_001_structure_tensor_ridge_heatmap.png",
        "002_ridge": candidate / "custom_part_scratch_002_structure_tensor_ridge_heatmap.png",
    }
    for key, source in raw_heatmaps.items():
        target = ASSET_DIR / f"{key}.png"
        colorize_heatmap(source, target)
        assets[key] = target
    return assets


def load_auto_mask_records() -> dict[str, dict]:
    path = MASK_ROOT / "metadata.jsonl"
    records = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        row = json.loads(line)
        records[Path(row["image_path"]).stem] = row
    return records


def load_ablation_rows() -> list[dict[str, str]]:
    path = REPORT_ROOT / "phase10_mask_quality" / "qwen" / "mask_quality_metrics.csv"
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def draw_bbox_asset(image_path: Path, region: list[int], output_path: Path, label: str) -> None:
    image = Image.open(image_path).convert("RGB")
    draw = ImageDraw.Draw(image)
    x1, y1, x2, y2 = map(int, region)
    width = max(3, image.width // 180)
    draw.rectangle((x1, y1, x2, y2), outline=(250, 190, 35), width=width)
    font = ImageFont.load_default()
    box = draw.textbbox((0, 0), label, font=font)
    tw, th = box[2] - box[0], box[3] - box[1]
    label_y = max(0, y1 - th - 8)
    draw.rounded_rectangle((x1, label_y, min(image.width, x1 + tw + 12), label_y + th + 7), radius=3, fill=(15, 23, 42))
    draw.text((x1 + 6, label_y + 3), label, fill=(255, 255, 255), font=font)
    image.save(output_path)


def make_title_triptych(output_path: Path) -> None:
    paths = [DATA_ROOT / "test" / "scratch" / f"{sample}.png" for sample in ("000", "001", "002")]
    canvas = Image.new("RGB", (900, 1200), (15, 23, 42))
    draw = ImageDraw.Draw(canvas)
    for i, path in enumerate(paths):
        image = Image.open(path).convert("RGB")
        image.thumbnail((760, 315), Image.Resampling.LANCZOS)
        x = (900 - image.width) // 2
        y = 60 + i * 370
        canvas.paste(image, (x, y))
        draw.rounded_rectangle((x - 4, y - 4, x + image.width + 4, y + image.height + 4), radius=8, outline=(45, 212, 191), width=4)
        draw.text((58, y + 10), f"WOOD {i:03d}", fill=(255, 255, 255), font=ImageFont.load_default())
    canvas.save(output_path)


def colorize_heatmap(source: Path, target: Path) -> None:
    gray = np.asarray(Image.open(source).convert("L"), dtype=np.float32) / 255.0
    low = float(np.percentile(gray, 4))
    high = float(np.percentile(gray, 99))
    norm = np.zeros_like(gray) if high <= low else np.clip((gray - low) / (high - low), 0.0, 1.0)
    rgb = np.zeros((*norm.shape, 3), dtype=np.uint8)
    rgb[..., 0] = np.clip(norm * 255, 0, 255).astype(np.uint8)
    rgb[..., 1] = np.clip((1.0 - np.abs(norm - 0.55) * 1.8) * 210, 0, 210).astype(np.uint8)
    rgb[..., 2] = np.clip((1.0 - norm) * 180, 0, 180).astype(np.uint8)
    Image.fromarray(rgb, mode="RGB").save(target)


def add_text(
    slide: object,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    size: float,
    color: RGBColor,
    *,
    bold: bool = False,
    italic: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
    font: str = FONT,
) -> object:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box


def add_bullet_list(
    slide: object,
    x: float,
    y: float,
    w: float,
    h: float,
    items: Iterable[str],
    size: float,
    color: RGBColor,
    *,
    bullet_color: RGBColor = TEAL,
    spacing: float = 4,
) -> object:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for i, item in enumerate(items):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.level = 0
        p.space_after = Pt(spacing)
        p.text = f"●  {item}"
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
        if p.runs:
            p.runs[0].font.color.rgb = color
    return box


def add_rect(
    slide: object,
    x: float,
    y: float,
    w: float,
    h: float,
    fill: RGBColor,
    *,
    line: RGBColor | None = None,
    radius: bool = False,
) -> object:
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.8)
    return shape


def add_line(slide: object, x1: float, y1: float, x2: float, y2: float, color: RGBColor, width: float = 1.2) -> object:
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(width)
    return line


def add_arrow(slide: object, x1: float, y1: float, x2: float, y2: float, color: RGBColor) -> object:
    arrow = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    arrow.line.color.rgb = color
    arrow.line.width = Pt(2.1)
    arrow.line.end_arrowhead = True
    return arrow


def add_pill(slide: object, x: float, y: float, w: float, h: float, text: str, color: RGBColor, font_size: float = 10) -> None:
    add_rect(slide, x, y, w, h, color, line=None, radius=True)
    add_text(slide, x + 0.04, y + 0.06, w - 0.08, h - 0.1, text, font_size, WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def add_callout(slide: object, x: float, y: float, w: float, h: float, title: str, body: str, color: RGBColor) -> None:
    add_rect(slide, x, y, w, h, WHITE, line=RGBColor(203, 213, 225), radius=True)
    add_rect(slide, x, y, 0.12, h, color, line=None)
    add_text(slide, x + 0.28, y + 0.16, w - 0.48, 0.28, title, 14, color, bold=True)
    add_text(slide, x + 0.28, y + 0.52, w - 0.48, h - 0.6, body, 10.5, NAVY)


def add_metric_card(slide: object, x: float, y: float, w: float, h: float, big: str, title: str, sub: str, color: RGBColor) -> None:
    add_rect(slide, x, y, w, h, WHITE, line=RGBColor(203, 213, 225), radius=True)
    add_rect(slide, x, y, 0.12, h, color, line=None)
    add_text(slide, x + 0.28, y + 0.16, w - 0.45, 0.48, big, 25, color, bold=True)
    add_text(slide, x + 0.28, y + 0.68, w - 0.45, 0.28, title, 12, NAVY, bold=True)
    add_text(slide, x + 0.28, y + 1.03, w - 0.45, 0.22, sub, 8.5, MUTED)


def draw_pipeline_box(slide: object, x: float, y: float, w: float, h: float, title: str, body: str, color: RGBColor, icon: str) -> None:
    add_rect(slide, x, y, w, h, WHITE, line=RGBColor(203, 213, 225), radius=True)
    add_pill(slide, x + 0.15, y + 0.15, 0.45, 0.38, icon, color, font_size=11)
    add_text(slide, x + 0.7, y + 0.17, w - 0.85, 0.3, title, 14, color, bold=True)
    add_text(slide, x + 0.2, y + 0.68, w - 0.4, 0.3, body, 11, NAVY, align=PP_ALIGN.CENTER)


def draw_stage_card(slide: object, x: float, y: float, w: float, h: float, title: str, body: str, color: RGBColor, number: int) -> None:
    add_rect(slide, x, y, w, h, WHITE, line=RGBColor(203, 213, 225), radius=True)
    add_pill(slide, x + 0.15, y + 0.15, 0.42, 0.36, str(number), color, font_size=10)
    add_text(slide, x + 0.66, y + 0.14, w - 0.82, 0.32, title, 14, color, bold=True)
    add_text(slide, x + 0.18, y + 0.62, w - 0.36, h - 0.72, body, 11, NAVY, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def draw_vertical_flow(slide: object, x: float, y: float, items: list[tuple[str, RGBColor]], arrow_color: RGBColor) -> None:
    for i, (label, color) in enumerate(items):
        yy = y + i * 0.91
        add_rect(slide, x, yy, 5.25, 0.64, WHITE, line=RGBColor(203, 213, 225), radius=True)
        add_rect(slide, x, yy, 0.1, 0.64, color, line=None)
        add_text(slide, x + 0.25, yy + 0.15, 4.75, 0.3, label, 13, NAVY, bold=True, align=PP_ALIGN.CENTER)
        if i < len(items) - 1:
            add_arrow(slide, x + 2.62, yy + 0.66, x + 2.62, yy + 0.87, arrow_color)


def add_factor_card(slide: object, x: float, y: float, w: float, h: float, title: str, desc: str, color: RGBColor) -> None:
    add_rect(slide, x, y, w, h, WHITE, line=RGBColor(203, 213, 225), radius=True)
    add_rect(slide, x, y, 0.11, h, color, line=None)
    add_text(slide, x + 0.28, y + 0.18, w - 0.48, 0.3, title, 14, color, bold=True)
    add_text(slide, x + 0.28, y + 0.58, w - 0.48, 0.46, desc, 10, NAVY)


def add_result_row(
    slide: object,
    x: float,
    y: float,
    w: float,
    h: float,
    sample: str,
    candidate: str,
    morphology: str,
    mask: str,
    color: RGBColor,
    compact: bool = False,
) -> None:
    add_rect(slide, x, y, w, h, PALE, line=RGBColor(226, 232, 240), radius=True)
    add_pill(slide, x + 0.12, y + 0.14, 0.58, 0.34, sample, color, font_size=10)
    size = 9 if compact else 10
    add_text(slide, x + 0.82, y + 0.1, w - 0.95, 0.22, candidate, size, NAVY, bold=True)
    add_text(slide, x + 0.82, y + 0.36, w - 0.95, 0.22, f"{morphology} → {mask}", size - 1, MUTED)


def add_table(
    slide: object,
    x: float,
    y: float,
    w: float,
    h: float,
    columns: list[str],
    rows: list[list[str]],
    col_widths: list[float],
) -> None:
    table_shape = slide.shapes.add_table(len(rows) + 1, len(columns), Inches(x), Inches(y), Inches(w), Inches(h))
    table = table_shape.table
    total = sum(col_widths)
    for idx, width in enumerate(col_widths):
        table.columns[idx].width = Inches(w * width / total)
    for c, label in enumerate(columns):
        cell = table.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        cell.text = label
        format_cell(cell, 10, WHITE, bold=True, align=PP_ALIGN.CENTER)
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r % 2 else PALE
            cell.text = value
            color = GREEN if value == "PASS" else NAVY
            format_cell(cell, 9.5, color, bold=value == "PASS", align=PP_ALIGN.CENTER)


def format_cell(cell: object, size: float, color: RGBColor, *, bold: bool = False, align: PP_ALIGN = PP_ALIGN.LEFT) -> None:
    frame = cell.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for paragraph in frame.paragraphs:
        paragraph.alignment = align
        for run in paragraph.runs:
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color


def add_picture_contain(slide: object, path: Path, x: float, y: float, w: float, h: float) -> object:
    with Image.open(path) as image:
        iw, ih = image.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    return slide.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))


def add_picture_cover(slide: object, path: Path, x: float, y: float, w: float, h: float, radius: bool = False) -> object:
    if radius:
        add_rect(slide, x - 0.02, y - 0.02, w + 0.04, h + 0.04, WHITE, line=RGBColor(203, 213, 225), radius=True)
    with Image.open(path) as image:
        iw, ih = image.size
    image_ratio = iw / ih
    box_ratio = w / h
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    if image_ratio > box_ratio:
        crop = (iw - ih * box_ratio) / (2 * iw)
        pic.crop_left = crop
        pic.crop_right = pic.crop_left
    else:
        crop = (ih - iw / box_ratio) / (2 * ih)
        pic.crop_top = crop
        pic.crop_bottom = pic.crop_top
    return pic


def add_citation(slide: object, text: str) -> None:
    add_text(slide, 0.72, 6.83, 11.9, 0.2, text, 7.5, MUTED, italic=True, align=PP_ALIGN.CENTER)


def write_outline() -> None:
    titles = [
        "Automatic Pseudo-Mask Generation",
        "Why automatic masks matter",
        "Research gap and project objective",
        "Current progress at a glance",
        "Auto-mask generation architecture",
        "User input contract",
        "Qwen localization",
        "Candidate bank",
        "Candidate scoring and selection",
        "Case 000: multi-scuff",
        "Cases 001/002: line scratches",
        "Morphology-aware label policy",
        "Purpose-specific mask outputs",
        "Uncertainty-aware learning",
        "Current wood result",
        "Candidate comparison",
        "Mask-policy ablation",
        "Manifest validation",
        "Downstream integration",
        "Advantages and limitations",
        "Next research upgrades",
        "Key takeaways",
        "Selected references",
    ]
    lines = [
        "# Automatic Pseudo-Mask Generation Presentation",
        "",
        f"PPTX: `{PPTX_PATH}`",
        "",
        "## Slide Outline",
        "",
    ]
    lines.extend(f"{i}. {title}" for i, title in enumerate(titles, start=1))
    lines.extend(
        [
            "",
            "## Early-Page Motivation Statement",
            "",
            "The deck explicitly states that many controllable industrial defect-generation methods require a user-supplied bounding box or pixel mask before synthesis, while noting that this is prevalent rather than universal.",
            "",
            "## Main Current Result",
            "",
            "```text",
            "000 -> patchcore_guided -> multi_scuff -> soft_mask_only -> training_soft",
            "001 -> ensemble_consensus -> scratch_band -> hard_mask_ok -> training_medium",
            "002 -> ensemble_consensus -> scratch_band -> hard_mask_ok -> training_medium",
            "```",
        ]
    )
    OUTLINE_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")


def validate_presentation(path: Path) -> None:
    prs = Presentation(path)
    if len(prs.slides) != 23:
        raise RuntimeError(f"Expected 23 slides, found {len(prs.slides)}")
    for index, slide in enumerate(prs.slides, start=1):
        if not slide.shapes:
            raise RuntimeError(f"Slide {index} is empty")
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0:
                raise RuntimeError(f"Slide {index} has a shape outside the top/left boundary")
            if shape.left + shape.width > prs.slide_width + Inches(0.05):
                raise RuntimeError(f"Slide {index} has a shape beyond the right boundary")
            if shape.top + shape.height > prs.slide_height + Inches(0.05):
                raise RuntimeError(f"Slide {index} has a shape beyond the bottom boundary")


if __name__ == "__main__":
    main()
